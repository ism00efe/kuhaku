"""DOCX content python-docx's high-level API hides: nested tables, content controls,
merged cells, toolbar numbering, headers and footers."""

from __future__ import annotations

import io

import pytest

from kuhaku.tools.parsing import DocumentParser, ListBlock, Paragraph, Table

docx = pytest.importorskip("docx")
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _parse(document):
    return DocumentParser(ocr=None).parse(_save(document), filename="d.docx")


def test_nested_table_text_is_kept_in_its_cell() -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "outer"
    inner = table.rows[0].cells[1].add_table(rows=1, cols=2)
    inner.rows[0].cells[0].text = "in"
    inner.rows[0].cells[1].text = "ner"
    [parsed] = _parse(document).tables
    assert parsed.rows == (("outer", "in ner"),)


def test_block_level_content_controls_are_read() -> None:
    document = docx.Document()
    document.add_paragraph("before")
    hidden = document.add_paragraph("inside a content control")
    sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    sdt.append(content)
    hidden._p.addprevious(sdt)
    content.append(hidden._p)
    document.add_paragraph("after")
    assert [b.text for b in _parse(document).blocks] == [
        "before",
        "inside a content control",
        "after",
    ]


def test_merged_cells_are_read_once() -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "wide"
    table.cell(0, 2).text = "c"
    table.cell(1, 0).text = "x"
    table.cell(1, 1).merge(table.cell(1, 2)).text = "also wide"
    [parsed] = _parse(document).tables
    assert parsed.rows == (("wide", "", "c"), ("x", "also wide", ""))

    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "tall"
    table.cell(0, 1).text = "a"
    table.cell(1, 1).text = "b"
    [parsed] = _parse(document).tables
    assert parsed.rows == (("tall", "a"), ("", "b"))


def test_toolbar_numbered_list_paragraphs_are_ordered() -> None:
    document = docx.Document()
    for text, num_id in (("one", "5"), ("two", "5")):
        paragraph = document.add_paragraph(text, style="List Paragraph")
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl, num = OxmlElement("w:ilvl"), OxmlElement("w:numId")
        ilvl.set(qn("w:val"), "0")
        num.set(
            qn("w:val"), num_id
        )  # numId 5 is a decimal list in the default template
        num_pr.extend([ilvl, num])
        p_pr.append(num_pr)
    [block] = _parse(document).blocks
    assert isinstance(block, ListBlock) and block.ordered
    assert [item.text for item in block.items] == ["one", "two"]


def test_headers_and_footers_are_separate_blocks_with_provenance() -> None:
    document = docx.Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "ACME Ltd -- Confidential"
    section.footer.paragraphs[0].text = "Page footer"
    document.add_paragraph("Body text")
    blocks = list(_parse(document).blocks)

    assert [(b.text, b.provenance.region) for b in blocks] == [
        ("ACME Ltd -- Confidential", "header"),
        ("Body text", None),
        ("Page footer", "footer"),
    ]
    assert all(isinstance(b, Paragraph) for b in blocks)


def test_style_numbered_lists_still_resolve_through_numbering() -> None:
    document = docx.Document()
    document.add_paragraph("a", style="List Number")
    document.add_paragraph("b", style="List Bullet")
    blocks = list(_parse(document).blocks)
    assert [(type(b), b.ordered) for b in blocks] == [
        (ListBlock, True),
        (ListBlock, False),
    ]
    assert not any(isinstance(b, Table) for b in blocks)
